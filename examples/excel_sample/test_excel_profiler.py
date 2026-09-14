"""
Test end-to-end de excel-profiler sobre el sample sucio.

Ejecuta las recetas de skills/excel-profiler/SKILL.md y verifica:
- Detecta "Data" como hoja principal (no Portada ni Metadata).
- Detecta fila 5 como fila de headers.
- Reporta metadata previa (logo + subtítulo).
- Reporta merged cells relevantes.
- Carga el dataframe con headers limpios.
- Detecta columna email con 50% nulos (señal de alerta).
- Detecta columna pais como constante.
- Detecta fila TOTAL al final.
- Output: <basename>_profile.md en examples/excel_sample/.
"""
import sys
from pathlib import Path
import openpyxl
import pandas as pd
import numpy as np

EXAMPLE_DIR = Path(__file__).parent
SAMPLE = EXAMPLE_DIR / "ventas_q2_2026_dirty.xlsx"

# --- Recetas de la skill (copiadas de SKILL.md para el test) ---

def detect_sheet_with_data(sheets_info, sample_rows=50):
    """Score = % de filas con >=3 celdas no-vacías (filas "de datos")."""
    best_name, best_score = None, 0.0
    scores = {}
    for name, ws in sheets_info:
        rows_with_data = 0
        total_rows = 0
        for row in ws.iter_rows(min_row=1, max_row=sample_rows,
                                 values_only=True):
            total_rows += 1
            non_empty = sum(1 for c in row if c is not None and str(c).strip() != "")
            if non_empty >= 3:
                rows_with_data += 1
        score = rows_with_data / total_rows if total_rows > 0 else 0
        scores[name] = round(score, 3)
        if score > best_score and score > 0.5:
            best_name, best_score = name, score
    return best_name, scores


def detect_header_row(ws, max_scan=10):
    """Encuentra la fila donde están los headers reales.

    Estrategia de 2 pases:
    1. Primero busca una fila con todos los valores siendo strings no-numéricos
       y >=3 celdas no-vacías. Esto es el caso típico de headers.
    2. Si no encuentra, fallback a fila con tipos heterogéneos (>=3 celdas,
       mezcla de tipos) que NO esté cubierta por merged cells.
    """
    merged_ranges = list(ws.merged_cells.ranges)
    merged_rows = set()
    for mr in merged_ranges:
        for row_idx in range(mr.min_row, mr.max_row + 1):
            merged_rows.add(row_idx)

    # Pase 1: fila con todos strings (típico header)
    for row_idx in range(1, max_scan + 1):
        cells = [c.value for c in ws[row_idx]]
        non_empty = [c for c in cells if c is not None and str(c).strip() != ""]
        if len(non_empty) >= 3 and row_idx not in merged_rows:
            # ¿Todos los valores son strings no-numéricos?
            all_stringy = all(
                isinstance(c, str) and not c.replace(".", "").replace("-", "").isdigit()
                for c in non_empty
            )
            if all_stringy:
                return row_idx

    # Pase 2: fallback — fila con tipos heterogéneos
    for row_idx in range(1, max_scan + 1):
        cells = [c.value for c in ws[row_idx]]
        non_empty = [c for c in cells if c is not None and str(c).strip() != ""]
        if len(non_empty) >= 3 and row_idx not in merged_rows:
            types = {type(c).__name__ for c in non_empty}
            if len(types) >= 2:
                return row_idx

    # Pase 3: última opción — primera fila con >=3 no-vacías
    for row_idx in range(1, max_scan + 1):
        cells = [c.value for c in ws[row_idx]]
        non_empty = [c for c in cells if c is not None and str(c).strip() != ""]
        if len(non_empty) >= 3:
            return row_idx
    return 1


def merged_cells_in_columns(ws, header_row, n_cols):
    relevant = []
    for mr in ws.merged_cells.ranges:
        if mr.min_row < header_row and mr.max_col <= n_cols:
            relevant.append({
                "range": str(mr),
                "value": str(ws.cell(mr.min_row, mr.min_col).value),
                "spans_rows": mr.max_row - mr.min_row + 1,
                "spans_cols": mr.max_col - mr.min_col + 1,
            })
    return relevant


def find_total_rows(df):
    total_rows = []
    for idx, row in df.iterrows():
        for val in row.values:
            if isinstance(val, str) and "total" in val.lower():
                total_rows.append(idx)
                break
    return total_rows


# --- Test ---

print(f"📂 Sample: {SAMPLE}\n")

# 1. Carga estructural (read_only para eficiencia)
wb_ro = openpyxl.load_workbook(SAMPLE, read_only=True, data_only=True)
sheets_info = [(ws.title, ws) for ws in wb_ro.worksheets]
print(f"📋 Hojas encontradas: {[s[0] for s in sheets_info]}")

# 2. Selección de hoja
detected_sheet, scores = detect_sheet_with_data(sheets_info)
print(f"🎯 Hoja detectada con datos: {detected_sheet}")
print(f"   Scores: {scores}")
assert detected_sheet == "Data", f"Esperaba 'Data', obtuve '{detected_sheet}'"
print("   ✅ Hoja correcta\n")

# 3. Recargar en modo normal (read_only=False) para acceder a merged_cells
wb = openpyxl.load_workbook(SAMPLE, read_only=False, data_only=True)
ws = wb[detected_sheet]

# 4. Detección de header
header_row = detect_header_row(ws)
print(f"📍 Fila de headers detectada: {header_row}")
assert header_row == 5, f"Esperaba fila 5, obtuve {header_row}"
print("   ✅ Fila correcta\n")

# 4. Merged cells
n_cols = 6  # cliente_id, nombre, email, monto_q2, pais, segmento
merged = merged_cells_in_columns(ws, header_row, n_cols)
print(f"🔗 Merged cells relevantes ({len(merged)}):")
for mr in merged:
    print(f"   - {mr['range']}: {mr['value']!r}")
assert len(merged) >= 2, f"Esperaba >=2 merged cells, obtuve {len(merged)}"
print("   ✅ Merged cells reportadas\n")

# 5. Carga del dataframe
df = pd.read_excel(SAMPLE, sheet_name=detected_sheet, header=header_row - 1)
print(f"📊 DataFrame cargado: {df.shape[0]} filas × {df.shape[1]} columnas")
print(f"   Columnas: {list(df.columns)}")
expected_cols = ["cliente_id", "nombre", "email", "monto_q2", "pais", "segmento"]
assert list(df.columns) == expected_cols, f"Columnas inesperadas: {list(df.columns)}"
print("   ✅ Headers limpios\n")

# 6. Perfil + alertas
NULL_TOKENS = {"", "nan", "NaN", "NAN", "null", "NULL", "None", "none", "?"}
df = df.replace(NULL_TOKENS, np.nan)

nulls = df.isna().sum()
pct_nulls = (df.isna().mean() * 100).round(2)
print(f"⚠️  Nulos por columna:")
for col in df.columns:
    print(f"   - {col}: {nulls[col]} ({pct_nulls[col]}%)")

email_null_pct = pct_nulls["email"]
assert email_null_pct >= 30, f"Esperaba email con >=30% nulos, obtuve {email_null_pct}%"
print(f"   ✅ email con {email_null_pct}% nulos — alerta correcta\n")

pais_unique = df["pais"].nunique()
print(f"🔁 Valores únicos en 'pais': {pais_unique}")
assert pais_unique == 1, f"Esperaba 'pais' constante, obtuve {pais_unique} únicos"
print("   ✅ pais constante — alerta correcta\n")

# 7. Fila TOTAL
total_rows = find_total_rows(df)
print(f"📍 Filas con 'Total': {total_rows}")
assert len(total_rows) >= 1, "Esperaba >=1 fila con 'Total'"
print("   ✅ Fila TOTAL detectada\n")

# 8. Output markdown
out_path = EXAMPLE_DIR / "ventas_q2_2026_dirty_profile.md"
with open(out_path, "w") as f:
    f.write(f"# Perfil: ventas_q2_2026_dirty.xlsx\n\n")
    f.write("## Metadata del workbook\n\n")
    f.write(f"- Hojas: {len(sheets_info)}\n")
    f.write(f"- Hoja detectada con datos: **{detected_sheet}**\n")
    f.write(f"- Fila de headers detectada: **{header_row}**\n")
    f.write(f"- Tamaño: {SAMPLE.stat().st_size / 1024:.2f} KB\n\n")

    f.write("## Metadata previa al header\n\n")
    for row_idx in range(1, header_row):
        cells = [(c.coordinate, c.value) for c in ws[row_idx] if c.value is not None]
        if cells:
            f.write(f"- Fila {row_idx}: {cells}\n")

    f.write("\n## Merged cells relevantes\n\n")
    for mr in merged:
        f.write(f"- `{mr['range']}`: {mr['value']!r}\n")

    f.write("\n## Perfil por columna\n\n")
    f.write("| columna | n_nulos | pct_nulos | n_unicos |\n")
    f.write("|---|---|---|---|\n")
    for col in df.columns:
        f.write(f"| {col} | {nulls[col]} | {pct_nulls[col]} | {df[col].nunique()} |\n")

    f.write("\n## Tabla completa\n\n")
    f.write(f"- filas: {len(df)}\n- columnas: {df.shape[1]}\n")
    f.write(f"- duplicados: {df.duplicated().sum()}\n")

    f.write("\n## Señales de alerta\n\n")
    f.write(f"- ⚠️ Columna `email` con {email_null_pct}% de nulos\n")
    f.write(f"- ⚠️ Columna `pais` es constante (1 valor único: '{df['pais'].dropna().iloc[0] if df['pais'].notna().any() else 'N/A'}')\n")
    f.write(f"- ⚠️ Fila con 'Total' detectada en posición: {total_rows} (debe filtrarse antes de agregar)\n")

print(f"📝 Reporte guardado: {out_path}")
print("\n✅ Todos los tests pasaron.")