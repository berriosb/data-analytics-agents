"""
Genera un Excel "sucio" de muestra para test de excel-profiler.

Estructura:
- 3 hojas: "Portada", "Metadata", "Data"
- Hoja "Data": logo/título merged en filas 1-2, subtítulo en fila 3, fila vacía en 4,
  headers reales en fila 5, datos desde fila 6.
- 1 columna con >30% nulos ("email")
- 1 columna constante ("pais")
- 1 columna ID único ("cliente_id")
- Fila "Total" mal formateada al final.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

OUT = Path(__file__).parent / "ventas_q2_2026_dirty.xlsx"

wb = openpyxl.Workbook()

# Hoja 1: Portada (debe ser ignorada — score < 0.1)
ws_portada = wb.active
assert ws_portada is not None
ws_portada.title = "Portada"
ws_portada["A1"] = "REPORTE CONFIDENCIAL"
ws_portada["A1"].font = Font(size=20, bold=True)
ws_portada["A3"] = "Elaborado por: Equipo de Finanzas"
ws_portada["A4"] = "Fecha: 2026-07-01"
ws_portada["A5"] = "Clasificación: Interno"

# Hoja 2: Metadata (debe ser ignorada — baja densidad)
ws_meta = wb.create_sheet("Metadata")
ws_meta["A1"] = "Variables y definiciones"
ws_meta["A2"] = "cliente_id: ID único del cliente"
ws_meta["A3"] = "nombre: Razón social"
ws_meta["A4"] = "email: Email de contacto"
ws_meta["A5"] = "monto_q2: Venta en pesos CLP"
ws_meta["A6"] = "pais: País (constante: Chile)"

# Hoja 3: Data (esta debe ser detectada)
ws = wb.create_sheet("Data")

# Fila 1-2: logo merged (merged cells, fuente grande)
ws.merge_cells("A1:F2")
ws["A1"] = "VENTAS Q2 2026 — DETALLE POR CLIENTE"
ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

# Fila 3: subtítulo merged
ws.merge_cells("A3:F3")
ws["A3"] = "Fuente: ERP interno · Generado: 2026-07-01 · Moneda: CLP"
ws["A3"].font = Font(italic=True, size=10)
ws["A3"].alignment = Alignment(horizontal="center")

# Fila 4: vacía (gap)

# Fila 5: headers reales
headers = ["cliente_id", "nombre", "email", "monto_q2", "pais", "segmento"]
for col, h in enumerate(headers, start=1):
    cell = ws.cell(row=5, column=col, value=h)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="D9E1F2")

# Filas 6-25: datos
data = [
    ("C001", "Comercial Andina SpA", "contacto@andinacl", 45_200_000, "Chile", "Retail"),
    ("C002", "Distribuidora del Sur Ltda", None, 28_900_000, "Chile", "Retail"),
    ("C003", "Servicios Norte S.A.", "finanzas@norte.cl", 67_400_000, "Chile", "Servicios"),
    ("C004", "Importadora Austral SpA", None, 12_300_000, "Chile", "Retail"),
    ("C005", "Constructora Cordillera", "ops@cordilleracl", 89_100_000, "Chile", "Construcción"),
    ("C006", "Agroindustria del Maule", None, 34_500_000, "Chile", "Agro"),
    ("C007", "Tech Solutions Chile", "hola@techsolutions.cl", 56_700_000, "Chile", "Tech"),
    ("C008", "Logística Pacífico Ltda", None, 41_200_000, "Chile", "Logística"),
    ("C009", "Energía Renovable SpA", "gerencia@errnclr", 73_800_000, "Chile", "Energía"),
    ("C010", "Manufactura del Biobío", None, 52_400_000, "Chile", "Manufactura"),
    ("C011", "Retail Express S.A.", "compras@retailexpress.cl", 38_900_000, "Chile", "Retail"),
    ("C012", "Servicios Financieros Sur", None, 91_300_000, "Chile", "Finanzas"),
    ("C013", "Industrial Atacama SpA", "ventas@atacama-ind.cl", 27_600_000, "Chile", "Manufactura"),
    ("C014", "Comercial Valparaíso Ltda", None, 19_400_000, "Chile", "Retail"),
    ("C015", "Transporte Andino S.A.", "logistica@transandino.cl", 62_100_000, "Chile", "Logística"),
    ("C016", "Salud y Bienestar SpA", None, 48_700_000, "Chile", "Salud"),
    ("C017", "Educación Continua Ltda", "admision@educontinua.cl", 15_200_000, "Chile", "Educación"),
    ("C018", "Turismo Patagónico SpA", None, 33_800_000, "Chile", "Turismo"),
    ("C019", "Minería del Norte S.A.", "contacto@mineranorte.cl", 95_500_000, "Chile", "Minería"),
    ("C020", "Inversiones del Sur SpA", None, 71_600_000, "Chile", "Finanzas"),
]

for row_idx, row_data in enumerate(data, start=6):
    for col_idx, val in enumerate(row_data, start=1):
        ws.cell(row=row_idx, column=col_idx, value=val)

# Fila 26: TOTAL mal formateado (mezcla strings y números)
ws.cell(row=26, column=1, value="TOTAL")
ws.cell(row=26, column=1).font = Font(bold=True)
ws.cell(row=26, column=4, value="=SUM(D6:D25)")  # fórmula
ws.cell(row=26, column=4).number_format = '"$"#,##0'

# Anchos de columna
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 32
ws.column_dimensions["C"].width = 32
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 8
ws.column_dimensions["F"].width = 14

wb.save(OUT)
print(f"✅ Generado: {OUT}")
print(f"   - 3 hojas (Portada debe ser ignorada)")
print(f"   - Hoja Data: logo merged en 1-2, subtítulo en 3, headers en fila 5")
print(f"   - 20 filas de datos con email 50% nulo, pais constante")
print(f"   - Fila TOTAL mal formateada al final")